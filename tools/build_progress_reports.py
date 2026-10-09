"""Tạo hai báo cáo tiến độ bằng số liệu hiện có; không huấn luyện mô hình.

    python tools/build_progress_reports.py
    python tools/build_progress_reports.py --pdf

Chạy lại sau khi cập nhật summary/artifacts để ghi nhận trạng thái mới.
Báo cáo là bản chuẩn bị nộp, không xác nhận đã nộp hoặc lùi ngày báo cáo.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

try:
    from tools.build_report_docx import (
        add_reference_table, add_table, configure, export_pdf_with_word,
        inline, page_footer, set_font, set_page, table_rows,
    )
except ModuleNotFoundError:
    # Cho phép gọi trực tiếp file trong tools/ mà không sửa sys.path.
    from build_report_docx import (
        add_reference_table, add_table, configure, export_pdf_with_word,
        inline, page_footer, set_font, set_page, table_rows,
    )


ROOT = Path(__file__).resolve().parents[1]
SEEDS = (42, 123, 2026)
ARCHITECTURES = {
    "bert": ("C1 BERT-base-cased", "Quốc Khánh"),
    "roberta": ("C2 RoBERTa-base", "Đức Trí"),
    "distilbert": ("C3 DistilBERT-base-uncased", "Nhật Huy"),
}


def read_json(path, default=None):
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise RuntimeError(f"Tệp {path} đang cập nhật hoặc không hợp lệ; chạy lại sau khi ghi xong.") from error


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def score(value):
    return "—" if value is None else f"{float(value):.4f}"


def md_table(headers, rows):
    def safe(value):
        return str(value).replace("|", "/").replace("\n", " ")
    return "\n".join([
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
        *("| " + " | ".join(safe(x) for x in row) + " |" for row in rows),
    ])


def snapshot():
    """Đọc snapshot summary, bổ sung trạng thái run hiện tại từ metadata."""
    summary_path = ROOT / "reports/project_results/summary.json"
    summary = read_json(summary_path, {"records": [], "averages": [], "missing": [], "complete": False})
    baseline = read_json(ROOT / "reports/baseline_validation/summary.json", {"configurations": []})
    eda = read_json(ROOT / "reports/summary.json", {})
    run_rows = []
    for architecture, (label, owner) in ARCHITECTURES.items():
        for seed in SEEDS:
            directory = ROOT / f"data/processed/transformers/{architecture}/seed_{seed}/full/standard"
            metadata = read_json(directory / "run_metadata.json", {})
            metrics = read_json(directory / "validation_metrics.json", {})
            expected = metadata.get("artifact_sha256", {}).get("validation_metrics.json")
            hashes_match = not expected or ((directory / "validation_metrics.json").exists() and digest(directory / "validation_metrics.json") == expected)
            completed = bool(metadata.get("completed") and not metadata.get("smoke") and metrics.get("n_samples") == 5426 and metrics.get("n_labels") == 28 and hashes_match)
            if completed:
                state = "Full validation hoàn thành"
            elif metadata:
                state = "Có metadata; chưa hoàn thành full"
            else:
                state = "Chưa có artifacts full"
            run_rows.append({"architecture": architecture, "label": label, "owner": owner, "seed": seed,
                             "state": state, "completed": completed, "metrics": metrics if completed else {},
                             "metadata": metadata, "path": str(directory.relative_to(ROOT)).replace("\\", "/")})
    zero_shot = {}
    for mode in ("full", "smoke"):
        directory = ROOT / f"data/processed/zero_shot/{mode}"
        metadata = read_json(directory / "run_metadata.json", {})
        metrics = read_json(directory / "validation_metrics.json", {})
        completed = metadata.get("status") == "complete" and bool(metrics)
        if mode == "full":
            completed = completed and not metadata.get("smoke") and metrics.get("n_samples") == 5426
        if completed:
            zero_shot = {"mode": mode, "metadata": metadata, "metrics": metrics,
                         "path": str(directory.relative_to(ROOT)).replace("\\", "/")}
            break
    return {"summary": summary, "baseline": baseline, "eda": eda, "runs": run_rows,
            "zero_shot": zero_shot, "selection": read_json(ROOT / "data/processed/transformers/selected_model.json", {}),
            "demo": read_json(ROOT / "reports/demo_verification.json", {}),
            "demo_ui": read_json(ROOT / "reports/demo_ui/evidence.json", {}),
            "verification": read_json(ROOT / "reports/verification_project.json", {}),
            "summary_hash": digest(summary_path) if summary_path.exists() else None}


def roles_table(data):
    counts = {arch: sum(row["completed"] for row in data["runs"] if row["architecture"] == arch) for arch in ARCHITECTURES}
    b_status = "B full đã có" if data["zero_shot"].get("mode") == "full" else "B chưa đủ full"
    return md_table(["Thành viên", "Vai trò", "Bằng chứng hiện có", "% công sức"], [
        ["Bảo Duy Nguyễn", "A; điều phối B; data/metrics và bảng nâng cao", f"A full validation; {b_status}", "____________________"],
        ["Quốc Khánh", "C1 BERT; phần đầu/tổng hợp báo cáo", f"C1 {counts['bert']}/3 seed full; xem bảng run", "____________________"],
        ["Đức Trí", "C2 RoBERTa; hỗ trợ/bàn giao B", f"C2 {counts['roberta']}/3 seed full; xem bảng run", "____________________"],
        ["Nhật Huy", "C3 DistilBERT; tích hợp demo best C", f"C3 {counts['distilbert']}/3 seed full; demo theo hồ sơ", "____________________"],
    ])


def baseline_table(data):
    rows = []
    for item in data["baseline"].get("configurations", []):
        rows.append([item["configuration"], item.get("n_validation", item.get("n_samples", "—")),
                     score(item.get("macro_f1")), score(item.get("micro_f1")),
                     score(item.get("micro_precision")), score(item.get("micro_recall")), score(item.get("hamming_loss"))])
    return md_table(["Cấu hình A", "Val N", "Macro-F1", "Micro-F1", "Micro-P", "Micro-R", "Hamming"], rows) if rows else "Chưa có summary baseline đầy đủ tại thời điểm render."


def b_section(data):
    b = data["zero_shot"]
    if not b:
        return "B đã có hướng triển khai BART-MNLI/pipeline; chưa có số từ run hoàn thành tại thời điểm render. Phần này chưa đạt tiêu chí số liệu cụ thể của báo cáo tiến độ 1."
    metrics = b["metrics"]
    mode = "FULL validation" if b["mode"] == "full" else "SMOKE/PILOT; không thay benchmark full"
    text = md_table(["Phạm vi B", "N", "Ngưỡng", "Macro-F1", "Micro-F1", "Hamming"], [[
        mode, metrics.get("n_samples", "—"), metrics.get("threshold", 0.5),
        score(metrics.get("macro_f1")), score(metrics.get("micro_f1")), score(metrics.get("hamming_loss")),
    ]])
    return text + "\n\n" + (
        f"Nguồn: `{b['path']}/run_metadata.json` và `validation_metrics.json`. "
        f"Template: `{b['metadata'].get('hypothesis_template', 'Chưa ghi')}`. "
        "B không cập nhật trọng số trên GoEmotions; kết quả threshold hiệu chỉnh nếu có phải tách khỏi hàng gốc. "
        + ("Đã có số full; kiểm frozen protocol trước test." if b["mode"] == "full" else "Chưa có B full; điểm pilot này không được dùng kết luận B tốt/kém hơn A hoặc C.")
    )


def runs_table(data):
    return md_table(["Kiến trúc", "Seed", "Trạng thái", "Macro-F1 val", "Micro-F1 val", "Hamming val"], [
        [row["label"], row["seed"], row["state"], score(row["metrics"].get("macro_f1")),
         score(row["metrics"].get("micro_f1")), score(row["metrics"].get("hamming_loss"))]
        for row in data["runs"]
    ])


def environment_table(data):
    rows = []
    for architecture, (label, owner) in ARCHITECTURES.items():
        available = [row for row in data["runs"] if row["architecture"] == architecture and row["metadata"]]
        if not available:
            rows.append([label, owner, "Chưa có metadata full", "—"])
            continue
        metadata = available[0]["metadata"]
        env = metadata.get("environment", {})
        config = metadata.get("config", {})
        device = env.get("gpu", env.get("device", config.get("device", "—")))
        versions = f"Python {env.get('python', '—')}; torch {env.get('torch', '—')}; Transformers {env.get('transformers', '—')}"
        params = f"lr={config.get('learning_rate', '—')}; epochs={config.get('epochs', '—')}; batch={config.get('batch_size', '—')}; accum={config.get('gradient_accumulation', '—')}; max_length={config.get('max_length', '—')}; {config.get('padding', '—')}"
        rows.append([label, owner, str(device) + "; " + versions, params])
    return md_table(["Kiến trúc", "Phụ trách", "Môi trường thực tế", "Config từ metadata"], rows)


def mean_std_table(data):
    rows = []
    for architecture, (label, _) in ARCHITECTURES.items():
        group = [row for row in data["runs"] if row["architecture"] == architecture and row["completed"]]
        if len(group) < 3:
            rows.append([label, len(group), "Chưa đủ 3 seed", "Chưa đủ 3 seed", "Chưa xếp hạng"])
            continue
        values = {}
        for name in ("macro_f1", "micro_f1", "hamming_loss"):
            nums = [row["metrics"][name] for row in group]
            mean = sum(nums) / len(nums)
            std = math.sqrt(sum((x - mean) ** 2 for x in nums) / (len(nums) - 1))
            values[name] = f"{mean:.4f} ± {std:.4f}"
        rows.append([label, len(group), values["macro_f1"], values["micro_f1"], values["hamming_loss"]])
    return md_table(["Kiến trúc", "Full seeds", "Macro-F1 mean±std val", "Micro-F1 mean±std val", "Hamming mean±std val"], rows)


def append_references(original_ids):
    all_refs = {item["id"]: item for item in read_json(ROOT / "reports/references_ieee.json", [])}
    return "\n\n".join(f"[{index}] {all_refs[source]['ieee']}" for index, source in enumerate(original_ids, 1))


def progress_one(data, stamp):
    completed = sum(row["completed"] for row in data["runs"])
    next_steps = ("- Đã đủ B full và 9/9 run C; duy trì logs/checkpoints/scores/config khi bàn giao.\n"
                  "- Đã khóa lựa chọn/ngưỡng trên validation và có test; không chọn lại bằng test.\n"
                  "- Đã có ba case đối chiếu và hồ sơ demo; nhóm cần tự đọc, diễn giải và kiểm khi chuyển máy."
                  if data["summary"].get("complete") else
                  "- Hoàn tất B full và ba C đủ seed; giữ logs/checkpoints/scores/config.\n"
                  "- Dùng validation để khóa kiến trúc/seed/ngưỡng trước test.\n"
                  "- Đối chiếu ít nhất ba nhóm lỗi cùng ID C1/C2/C3; dựng demo từ C được chọn.")
    return f"""# BÁO CÁO TIẾN ĐỘ LẦN 1

## Phân loại cảm xúc đa nhãn với GoEmotions

**Ngày cập nhật:** {stamp}. Đây là bản chuẩn bị báo cáo theo hiện trạng, không xác nhận đã nộp giảng viên.

**Giảng viên / mã lớp / năm học, học kỳ:** ____________________

**MSSV bốn thành viên:** ____________________

## 1. Bài nền tảng và mục tiêu

Demszky và cs. công bố GoEmotions tại ACL 2020: bình luận Reddit tiếng Anh, 27 cảm xúc cùng neutral, có thể nhiều nhãn trong một câu; baseline chính của bài dùng BERT [1]. Đề tài giữ đủ taxonomy và mở rộng so sánh A cổ điển, B zero-shot, ba C fine-tune theo yêu cầu môn học. TF-IDF + Logistic Regression là phần cổ điển cô giao, không phải baseline BERT trong paper.

Mục tiêu là khảo sát dữ liệu, tạo scores đa nhãn, đánh giá công bằng và xây demo từ C được chọn bằng validation. Nâng cao bằng weighting/ngưỡng riêng phải có số trước/sau và nhãn hiếm. Dữ liệu tiếng Anh Reddit không tự chứng minh chất lượng cho tiếng Việt hoặc doanh nghiệp.

## 2. Dữ liệu đã khám phá

Official split gồm 43.410 train, 5.426 validation, 5.427 test; tổng 54.263 ở bản filtered/simplified. Bản thô có 58.009 bình luận [2]. Revision nhóm: `add492243ff905527e67aeb8b80c082af02207c3`.

EDA từ `reports/summary.json`: {data['eda'].get('label_assignments', '—')} lượt nhãn, {data['eda'].get('multi_label_samples', '—')} câu đa nhãn; tỷ lệ đa nhãn khoảng {float(data['eda'].get('multi_label_pct', 0)):.2f}%. Tỷ số support nhãn phổ biến/hiếm khoảng {float(data['eda'].get('train_imbalance_ratio', 0)):.2f}. Năm nhãn hiếm được chọn bằng train: grief 77, pride 111, relief 153, nervousness 164, embarrassment 303.

Ví dụ dữ liệu validation: ID `eczdvun`, “Thank you. I really appreciate your response”, nhãn admiration và gratitude; đây là annotation thật được dùng trong hồ sơ lỗi A. Label mapping/multi-hot có 28 cột; neutral được giữ theo nguồn. ID giữa split không trùng nhưng có văn bản trùng; benchmark chính giữ official split và ghi giới hạn này.

## 3. Kết quả baseline A

A dùng unigram/bigram TF-IDF fit train và 28 bộ LR OvR [3]. Module metric giữ đủ nhãn, `zero_division=0`; báo F1 micro/macro cùng P/R/Hamming [4]. Bảng sau là full validation, không phải test.

{baseline_table(data)}

Weighting thường tăng recall cùng nhãn thừa; tuned-val được đo trên chính validation chọn ngưỡng nên có thể lạc quan. Không dùng Macro-F1 0,4901 tuned-val để tuyên bố vượt F1 test của paper. Tệp nguồn: `reports/baseline_validation/summary.json` và artifacts model/scores/thresholds của từng biến thể.

## 4. Zero-shot B: kết quả và trạng thái thực tế

Checkpoint BART-large-MNLI được dùng qua pipeline, `multi_label=True`, đủ 28 candidate labels; model đã học MNLI trước đó nhưng nhóm không fine-tune trên GoEmotions [5].

{b_section(data)}

## 5. Tiến độ ba kiến trúc và môi trường

Tại thời điểm cập nhật có {completed}/9 seed full được kiểm bằng metadata và validation metrics. Seed đã có metadata nhưng chưa hoàn thành không được tính là run full; smoke/pilot được tách khỏi bảng.

{runs_table(data)}

{environment_table(data)}

BERT cased theo lựa chọn kho GoEmotions; RoBERTa và DistilBERT có tokenizer/head riêng. Mỗi kiến trúc cần ba seed cùng cấu hình. Mốc mean±std chỉ tổng hợp khi đủ run; không tự báo std=0 với một seed. Tài liệu kiến trúc giải thích nền tảng C1/C2/C3 [6], [7], [8].

## 6. Phân công và mức hoàn thành theo artifacts

{roles_table(data)}

Đức Trí là “Thợ Săn Thập Cẩm”. B làm chung dưới điều phối của Duy; Trí vẫn phụ trách trọn C2. Huy tích hợp C thắng, không mặc định DistilBERT. Các tỷ lệ công sức để nhóm xác nhận, không tự chia đều.

## 7. Vướng mắc và đầu việc tiếp theo

{next_steps}
- Thiết bị từng bị ngủ trong quá trình C; thời gian elapsed có gián đoạn, không coi là benchmark tốc độ được kiểm soát.
- Cần đọc và giải thích code, xác nhận metadata hành chính và đóng góp trước khi nộp theo kênh cô chỉ định.

## 8. Hồ sơ kiểm chứng

Kho chung: https://github.com/trangkhanh-ai/goemotions-multilabel-classification

Summary tổng: `reports/project_results/summary.json`; SHA-256 tại thời điểm đọc: `{data['summary_hash'] or 'chưa có'}`. Script này chỉ đọc artifacts, không tự chạy GPU, không xác nhận đã nộp cô. Chạy lại sau khi cập nhật summary/runs để làm mới bản tiến độ.

## Tài liệu tham khảo

{append_references([1,2,5,16,13,8,9,10])}
"""


def progress_two(data, stamp):
    all_runs = sum(row["completed"] for row in data["runs"])
    summary = data["summary"]
    table = []
    for row in summary.get("averages", []):
        if row.get("threshold_mode") != "fixed":
            continue
        def mean(name):
            return score(row.get(name + "_mean")) + (" ± " + score(row[name + "_std"]) if row.get(name + "_std") is not None else "")
        table.append([row["system"], row["split"], row["n_runs"], mean("macro_f1"), mean("micro_f1"), mean("hamming_loss")])
    compare = md_table(["Hệ thống @0,5", "Split", "Run", "Macro-F1", "Micro-F1", "Hamming"], table) if table else "Chưa có summary so sánh."
    errors_path = ROOT / "reports/errors_test_standard_fixed/summary.md"
    errors = "Chưa có hồ sơ lỗi đối chiếu ba C trên test; phần này còn cần hoàn thành theo yêu cầu báo cáo tiến độ 2. Các ví dụ A dưới đây chỉ là bằng chứng đã có, không thay yêu cầu ba C."
    if errors_path.exists():
        errors = errors_path.read_text(encoding="utf-8")
        errors = re.sub(r"^# .*\n", "", errors, count=1)
        # Báo cáo tiến độ ngắn: giữ tối đa 12.000 ký tự, liên kết hồ sơ đầy đủ.
        if len(errors) > 12000:
            errors = errors[:12000].rsplit("\n", 1)[0] + "\n\nXem đầy đủ tại `reports/errors_test_standard_fixed/summary.md`."
    cases_path = ROOT / "reports/error_case_studies.md"
    if cases_path.exists():
        cases = re.sub(r"^# .*\n", "", cases_path.read_text(encoding="utf-8"), count=1)
        errors += "\n\n" + re.sub(r"^## ", "### ", cases, flags=re.MULTILINE)
    selection = data["selection"]
    if selection:
        selected_text = f"Selection manifest đã có: `{selection.get('architecture', '—')}`, seed `{selection.get('seed', '—')}`. Luật lựa chọn lấy từ validation; cần đọc `selected_model.json` và hồ sơ demo."
    else:
        selected_text = "Chưa có selection manifest best C. Không dùng A/B hoặc checkpoint smoke để thay demo môn học."
    if data["demo"]:
        evidence = data["demo"]
        demo_text = (f"Kiểm suy luận lúc {evidence.get('checked_at_utc', '—')}: "
                     f"{evidence.get('model_inference_status', '—')}, model {evidence.get('architecture', '—')}, "
                     f"seed {evidence.get('seed', '—')}; đối chiếu {len(evidence.get('validation_cases', []))} câu "
                     "với scores validation và kiểm luồng nhập. Đây là kiểm suy luận, không phải kiểm UI.")
    else:
        demo_text = "Chưa có hồ sơ kiểm demo tại thời điểm render. Khi hoàn tất cần URL chạy tại buổi demo hoặc ảnh/video, checkpoint/ngưỡng và so score với script suy luận."
    ui = data.get("demo_ui", {})
    if ui:
        demo_text += (f"\n\nKiểm UI riêng lúc {ui.get('checked_at_utc', '—')}: "
                      f"{ui.get('interface_status', '—')}, {ui.get('rendered_row_count', '—')}/28 hàng trong DOM. "
                      "Ảnh thực tế: reports/demo_ui/demo_ui.png; hash nằm trong evidence.json. "
                      "Hồ sơ ghi nhận thời điểm kiểm, không bảo đảm server đang chạy khi đọc báo cáo.")
    rare_path = ROOT / "reports/project_results/ANALYSIS.md"
    rare_text = "Chưa có bảng nhãn hiếm test đủ dữ liệu."
    if rare_path.exists():
        rare_text = rare_path.read_text(encoding="utf-8").split("## 2.", 1)[0]
        rare_text = re.sub(r"^# .*\n", "", rare_text, count=1)
        rare_text = re.sub(r"^## ", "### ", rare_text, flags=re.MULTILINE)
    verification = data.get("verification", {})
    final_tests = verification.get("final_unit_tests", {})
    notebook_tests = verification.get("notebook_verification", {})
    qa_text = (f"Kiểm mã: {final_tests.get('passed', '—')}/{final_tests.get('tests_run', '—')} unit tests đạt; "
               f"notebook {notebook_tests.get('passed', '—')}/{notebook_tests.get('completed', '—')} đạt. "
               "Nguồn verification_project.json; phép kiểm mã không thay benchmark full.")
    next_advanced = ("Đã có bảng test trước/sau với support, mean±std và cả mức giảm/không đổi. "
                     "C weighting chưa chạy; chín C standard dùng threshold tuning, không làm encoder học thêm. "
                     "Nhóm cần đọc và xác nhận các case, giữ nguyên protocol và artifacts khi bàn giao."
                     if summary.get("complete") else
                     "Đầu việc nâng cao tiếp theo: khóa threshold từng run, tổng hợp rare-label P/R/F1/support, "
                     "giữ cả nhãn không cải thiện và đánh giá cuối trên test. Nếu weighted C được bổ sung, phải train lại.")
    return f"""# BÁO CÁO TIẾN ĐỘ LẦN 2

## Phân loại cảm xúc đa nhãn với GoEmotions

**Ngày cập nhật:** {stamp}. Đây là bản chuẩn bị báo cáo hiện trạng, không xác nhận đã nộp giảng viên.

**Giảng viên / mã lớp / năm học, học kỳ:** ____________________

**MSSV bốn thành viên:** ____________________

## 1. Mục tiêu và phạm vi kiểm tra

Đề tài tiếp tục bài GoEmotions ACL 2020 [1], giữ 28 nhãn, split/mapping/metric dùng chung. Lần tiến độ 2 cần kết quả đủ ba kiến trúc × ba seed, mean±std so với A/B, ít nhất ba nhóm lỗi đối chiếu C, minh chứng demo và nâng cao. Các bảng được cập nhật bằng artifacts hiện có; nếu còn thiếu, báo trạng thái thay vì tạo số.

## 2. Kết quả A/B/C

{compare}

Macro/Micro-F1 và Hamming được đọc cùng Precision/Recall và support; Macro-F1 không phải tỷ lệ câu có cả tập nhãn đúng [2]. Summary có ghi split và luật ngưỡng; tuned-val dùng dữ liệu chọn ngưỡng không phải phép đánh giá độc lập.

### 2.1. Zero-shot B

Model BART-MNLI dùng trực tiếp, không cập nhật trọng số GoEmotions [3].

{b_section(data)}

### 2.2. Từng seed C

{runs_table(data)}

Hiện có {all_runs}/9 seed full theo metadata/scores. Bảng này chỉ tính full validation hợp lệ với 5.426 mẫu và 28 nhãn. Mỗi kiến trúc có ít nhất ba seed mới đủ mức nghiệm thu môn học.

### 2.3. Mean ± sample std

{mean_std_table(data)}

Sample std dùng mẫu số K−1. Một seed không cho sample std; chưa đủ ba seed thì chưa xếp hạng kiến trúc. Bảng từng seed và bảng tổng hợp phải được giữ cùng nhau; điểm checkpoint demo khác trung bình kiến trúc.

Ngưỡng riêng được chọn bằng validation của đúng model; không dùng test để tìm ngưỡng [4].

## 3. Phân tích lỗi có hệ thống

{errors}

Ba nhóm tự động gồm bỏ sót một phần nhãn, bỏ nhãn hiếm và cùng lúc FN/FP; các nhóm có thể chồng lấp. Ba case được diễn giải từ câu/nhãn/scores, không tự quy thành mỉa mai hoặc nguyên nhân trong encoder. Mỗi ví dụ đối chiếu giữ cùng ID và true labels. Một câu có thể đóng góp nhiều cặp FN/FP; không cộng mọi cặp để tính số câu sai.

Ví dụ A standard @0,5 đã kiểm: `eczwil0` bỏ sót pride; `ed832y6` dự đoán love thừa; `eczdvun` tìm gratitude nhưng thiếu admiration. Đây là ví dụ A thực tế, chưa phải lỗi của các C. Chỉ suy ra nguyên nhân sau khi đọc ngữ cảnh/annotation.

## 4. Demo D

{selected_text}

{demo_text}

{qa_text}

Gradio nối hàm suy luận với giao diện nhập văn bản [5]. Demo cần đúng checkpoint/tokenizer/threshold đã chọn, hỗ trợ nhiều nhãn, xử lý input rỗng và không ép neutral khi mọi score dưới ngưỡng. URL đang chạy/ảnh/video được điền sau khi xác minh, không tạo link giả.

**Link demo / minh chứng buổi trình bày:** ____________________

## 5. Nâng cao và nhãn hiếm

Đã có A standard/balanced và ngưỡng chung/riêng trên validation:

{baseline_table(data)}

Ngưỡng phải được chọn trên validation của đúng model, tránh dùng test để điều chỉnh [4]. Năm nhãn hiếm chọn bằng train: grief, pride, relief, nervousness, embarrassment. Có bảng F1 trước/sau A tại `reports/BASELINE_RESULTS.md` và bảng per-label. Cải tiến C nếu có phải báo cùng seed/config và mean±std đủ run; pilot một seed được ghi đúng pilot.

{rare_text}

{next_advanced}

## 6. Vai trò, đóng góp và phần chưa hoàn thành

{roles_table(data)}

Đức Trí là “Thợ Săn Thập Cẩm”; nhóm có bốn người. Tỷ lệ công sức để trống cho nhóm thống nhất bằng artifacts, không tự đánh đồng phân công và việc đã làm.

Theo summary hiện đọc: `complete={str(summary.get('complete', False)).lower()}`. Những phần cần bổ sung:

{chr(10).join('- ' + item for item in summary.get('missing', [])) or '- Không còn mục thiếu trong summary A/B/C; hồ sơ kiểm demo và đọc lỗi đã nêu trên. Nhóm cần xác nhận thông tin hành chính, đóng góp và kiểm demo khi chuyển máy.'}

Mức hoàn thành được ghi theo artifacts hiện có; không xác nhận điểm số hoặc đã nộp giảng viên. Các thông tin hành chính và tỷ lệ đóng góp vẫn để trống cho nhóm xác nhận.

## 7. Hồ sơ và bước hoàn tất

Summary: `reports/project_results/summary.json`; SHA-256: `{data['summary_hash'] or 'chưa có'}`. Cấu hình, môi trường và thời gian thực tế nằm trong `run_metadata.json`. Thời gian có gián đoạn do ngủ máy không được dùng như benchmark tốc độ được kiểm soát.

Kho chung: https://github.com/trangkhanh-ai/goemotions-multilabel-classification

Quy trình đã áp dụng: đủ seed → tổng hợp validation → chọn C/ngưỡng → khóa protocol → test → đối chiếu lỗi/demo → hoàn thiện báo cáo. Bước tiếp theo là đọc code, tập bảo vệ, điền thông tin hành chính và xác nhận đóng góp. Script tạo báo cáo chỉ đọc artifacts, không chạy GPU, không thay summary và không tự gửi/nộp.

## Tài liệu tham khảo

{append_references([1,16,13,19,25])}
"""


def render_docx(markdown, output):
    document = Document()
    configure(document)
    set_page(document.sections[0], "decimal", 1)
    page_footer(document.sections[0])
    document.core_properties.title = markdown.splitlines()[0].lstrip("# ")
    document.core_properties.author = "Bảo Duy Nguyễn; Quốc Khánh; Đức Trí; Nhật Huy"
    # Báo cáo tiến độ ngắn: tiêu đề/trang thông tin gọn trên trang đầu.
    lines = markdown.splitlines()
    i = 0
    in_references = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            title = heading[2]
            if i == 0:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(title)
                set_font(run, size=16)
                run.bold = True
            elif title == "Phân loại cảm xúc đa nhãn với GoEmotions":
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(title)
                set_font(run, size=14)
                run.bold = True
            else:
                document.add_heading(title, 2 if len(heading[1]) <= 2 else 3)
            in_references = title == "Tài liệu tham khảo"
            i += 1
            continue
        if in_references and re.match(r"^\[\d+\] ", line):
            refs = []
            while i < len(lines):
                m = re.match(r"^\[(\d+)\]\s+(.+)$", lines[i].strip())
                if m:
                    refs.append((m[1], m[2])); i += 1
                elif not lines[i].strip():
                    i += 1
                else:
                    break
            add_reference_table(document, refs)
            continue
        if line.startswith("|"):
            values = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values.append(lines[i]); i += 1
            add_table(document, table_rows(values))
            # Giữ header với ít nhất dòng dữ liệu đầu; tránh header nằm
            # một mình ở cuối trang khi bảng bắt đầu sát lề dưới.
            if document.tables:
                for cell in document.tables[-1].rows[0].cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
            continue
        if line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            inline(p, line[2:])
        elif line.startswith("<!--"):
            pass
        else:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(p, line)
        i += 1
    document.save(output)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", action="store_true", help="Xuất PDF bằng Word có sẵn; không cần để tạo MD/DOCX")
    args = parser.parse_args()
    stamp = datetime.now(timezone(timedelta(hours=7))).strftime("%d/%m/%Y %H:%M (Asia/Bangkok)")
    data = snapshot()
    result = []
    for number, text in ((1, progress_one(data, stamp)), (2, progress_two(data, stamp))):
        path = ROOT / f"reports/BAO_CAO_TIEN_DO_{number}.md"
        path.write_text(text, encoding="utf-8")
        docx = path.with_suffix(".docx")
        render_docx(text, docx)
        item = {"markdown": str(path), "docx": str(docx), "updated_at": stamp}
        if args.pdf:
            item["pdf"] = export_pdf_with_word(docx)
        result.append(item)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
