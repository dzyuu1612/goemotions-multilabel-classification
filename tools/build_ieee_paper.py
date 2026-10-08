"""Xuất bài đồ án riêng theo định dạng IEEE conference hai cột.

    python tools/build_ieee_paper.py
    python tools/build_ieee_paper.py --pdf

Chỉ đọc summary/CSV/metadata đã có; không tải hoặc huấn luyện mô hình.
Chạy lại sau khi summary được cập nhật để lấy số và trạng thái mới.
Bài sáu chương và bài hai cột là hai tài liệu riêng.
"""

from __future__ import annotations

import argparse
import csv
import ctypes
import hashlib
import json
import math
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

try:
    from tools.build_progress_reports import snapshot
    from tools.build_report_docx import export_pdf_with_word, hyperlink, inline, set_font, table_rows
except ModuleNotFoundError:
    from build_progress_reports import snapshot
    from build_report_docx import export_pdf_with_word, hyperlink, inline, set_font, table_rows


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports/BAI_BAO_GOEMOTIONS_IEEE_NOI_DUNG.md"
OUTPUT = ROOT / "reports/BAI_BAO_GOEMOTIONS_IEEE.docx"
COLUMN_CM = 3.5 * 2.54
SYSTEMS = ("A_standard", "A_balanced", "B_bart_mnli", "C_bert", "C_roberta", "C_distilbert")
SHORT = {"A_standard": "A-S", "A_balanced": "A-W", "B_bart_mnli": "B",
         "C_bert": "C1", "C_roberta": "C2", "C_distilbert": "C3"}
REF_IDS = (1, 2, 8, 9, 10, 11, 12, 13, 3, 5, 6, 7, 15, 19, 16, 17, 18, 23, 25, 26)
FORMAT_SOURCES = [
    {"url": "https://conferences.ieeeauthorcenter.ieee.org/write-your-paper/authoring-tools-and-templates/",
     "role": "Official IEEE Author Center links the conference templates."},
    {"url": "https://eit.r4.ieee.org/data/pdfs/IEEE-EIT-Paper_template.pdf",
     "role": "IEEE-hosted guide actually read: Letter margins, fonts, captions, citation order, no page numbers."},
    {"url": "https://ewh.ieee.org/soc/cas/dallas/dcas2015/DCAS2015_Author%20Kit.pdf",
     "role": "IEEE Author Kit actually read: 3.5-inch columns, 0.25-inch gap, 24/11/10/9/8-point sizes."},
    {"url": "https://www.ieee.org/content/dam/ieee-org/ieee/web/org/conferences/conference-template-letter.docx",
     "role": "Official download identified; returned empty HTTP 202 locally, so its contents were not verified."},
    {"url": "https://journals.ieeeauthorcenter.ieee.org/wp-content/uploads/sites/7/IEEE_Reference_Guide.pdf",
     "role": "Numeric IEEE references and source formatting checked separately."},
]


def read_csv(path):
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def number(value):
    if value in (None, "", "nan"):
        return None
    result = float(value)
    return result if math.isfinite(result) else None


def fmt(value, std=None, *, sign=False):
    value, std = number(value), number(std)
    if value is None:
        return "—"
    text = f"{value:+.4f}" if sign else f"{value:.4f}"
    return text if std is None else text + f" ± {std:.4f}"


def table(headers, rows):
    def safe(value):
        return str(value).replace("|", "/").replace("\n", " ")
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |",
                      *("| " + " | ".join(safe(v) for v in row) + " |" for row in rows)])


def aggregate_lookup(data):
    return {(row["system"], row["split"], row["threshold_mode"]): row
            for row in data["summary"].get("averages", [])}


def value(row, metric):
    return fmt(row.get(metric + "_mean"), row.get(metric + "_std")) if row else "—"


def abstract(data):
    completed = sum(run["completed"] for run in data["runs"])
    complete = data["summary"].get("complete", False)
    metrics = {row["configuration"]: row for row in data["baseline"].get("configurations", [])}
    standard, improved = metrics.get("standard_fixed", {}), metrics.get("balanced_tuned", {})
    status = ("Đã có kết quả full A/B/C; mỗi C gồm ba seed, mean và sample standard deviation. "
              if complete else
              f"Bản cập nhật này có {completed}/9 run C full; những bảng thiếu chưa được dùng xếp hạng. ")
    findings = ""
    if standard and improved:
        findings = (f"Trên validation, A standard @0,5 đạt Macro-F1 {fmt(standard.get('macro_f1'))}; "
                    f"A balanced với ngưỡng riêng đạt {fmt(improved.get('macro_f1'))}. "
                    "Điểm tuned-validation có thể lạc quan vì dùng lại dữ liệu chọn ngưỡng. ")
    return ("**Tóm tắt—** Nghiên cứu triển khai phân loại cảm xúc đa nhãn trên GoEmotions, "
            "gồm 27 cảm xúc và neutral, giữ official split 43.410/5.426/5.427 câu train/validation/test. "
            "A sử dụng TF-IDF + One-vs-Rest Logistic Regression; B dùng BART-MNLI zero-shot không fine-tune; "
            "C gồm BERT-base-cased, RoBERTa-base và DistilBERT với seed 42, 123, 2026. "
            "Nghiên cứu đánh giá Macro/Micro-F1, precision/recall và Hamming Loss; khảo sát class weighting, "
            "ngưỡng riêng và F1 năm nhãn hiếm xác định từ train. " + findings + status +
            "Mọi số lấy từ artifacts thực tế, tách smoke/full và validation/test. Demo được thiết kế dùng C chọn "
            "bằng validation. Bài viết diễn giải đầu ra, lỗi và giới hạn domain; liên hệ chuỗi dữ liệu→quyết định→giá trị "
            "trong Case Study 4 của Jay Lee, không quy số tiết kiệm của nhà máy thành ROI của mô hình NLP.")


def results(data, stamp):
    summary = data["summary"]
    averages = aggregate_lookup(data)
    completed = sum(run["completed"] for run in data["runs"])
    b_full = data["zero_shot"].get("mode") == "full"
    status = "Đủ hồ sơ benchmark theo summary" if summary.get("complete") else "Bản cập nhật chưa đủ benchmark cuối"
    parts = [f"**Trạng thái {stamp}: {status}.** Có {len(summary.get('records', []))}/72 hàng kết quả "
             f"và {len(summary.get('averages', []))}/36 nhóm tổng hợp theo thiết kế; C full {completed}/9, "
             f"B full {'đã có' if b_full else 'chưa có'}. Số hàng phụ thuộc artifact đã hoàn thành, không tính smoke. "
             "Dấu — là thiếu/chưa đủ ba seed hoặc không áp dụng, không phải F1=0. A-S là A standard; "
             "A-W là A balanced; C1/C2/C3 là BERT/RoBERTa/DistilBERT. Val là validation, N=5.426; Test N=5.427."]

    fixed_rows = []
    for system in SYSTEMS:
        validation = averages.get((system, "validation", "fixed"))
        test = averages.get((system, "test", "fixed"))
        # Không dùng một seed để tạo hàng mean/std của một kiến trúc C.
        count = validation.get("n_runs", "—") if validation else "—"
        fixed_rows.append([SHORT[system], count, value(validation, "macro_f1"), value(test, "macro_f1"),
                           value(test, "micro_f1"), value(test, "hamming_loss")])
    parts += ["**BẢNG I. SO SÁNH HỆ THỐNG GỐC, NGƯỠNG 0,5.**",
              table(["Hệ", "Run", "Val Macro", "Test Macro", "Test Micro", "Test H"], fixed_rows),
              "Cần đủ ba seed để có hàng C mean±std. Một C đã hoàn tất vẫn xuất hiện trong bảng từng seed ở dưới; "
              "không điền điểm đó vào hàng trung bình ba seed."]

    a_rows = []
    for row in data["baseline"].get("configurations", []):
        mode = row["configuration"].replace("standard", "A-S").replace("balanced", "A-W")
        a_rows.append([mode, fmt(row.get("macro_f1")), fmt(row.get("micro_f1")),
                       fmt(row.get("micro_precision")), fmt(row.get("micro_recall")), fmt(row.get("hamming_loss"))])
    if a_rows:
        parts += ["**BẢNG II. ABLATION A TRÊN VALIDATION.**",
                  table(["Cấu hình", "Macro", "Micro", "Pμ", "Rμ", "H"], a_rows),
                  "fixed: ngưỡng 0,5; global: một ngưỡng chọn trên val; tuned: riêng từng nhãn chọn trên val. "
                  "Các số là full validation, nhưng global/tuned không phải ước lượng độc lập với bước calibration."]

    test_rows = []
    for system in SYSTEMS:
        group = {mode: averages.get((system, "test", mode)) for mode in ("fixed", "global", "tuned")}
        if not any(group.values()):
            continue
        before, after = group["fixed"], group["tuned"]
        difference = after["macro_f1_mean"] - before["macro_f1_mean"] if before and after else None
        test_rows.append([SHORT[system], value(before, "macro_f1"), value(group["global"], "macro_f1"),
                          value(after, "macro_f1"), value(after, "micro_f1"), fmt(difference, sign=True)])
    if test_rows:
        parts += ["**BẢNG III. ABLATION NGƯỠNG TRÊN TEST ĐÃ KHÓA.**",
                  table(["Hệ", "MF fixed", "MF global", "MF tuned", "Micro tuned", "Δ MF"], test_rows),
                  "MF là Macro-F1; Δ của bảng này so fixed→tuned trên cùng hệ. Std, nếu có, là giữa seed; "
                  "không tuning lại bằng test. Hamming/P/R cho mọi chế độ vẫn được giữ trong mean_std.csv."]
    else:
        parts += ["**Bảng test ablation chưa có đủ artifact.** Chưa thể kết luận cải thiện tổng quát trên test từ bảng tuned-validation."]

    pr_rows = []
    for system in SYSTEMS:
        split = "test" if (system, "test", "fixed") in averages else "validation"
        row = averages.get((system, split, "fixed"))
        if row:
            pr_rows.append([SHORT[system], "Test" if split == "test" else "Val", value(row, "macro_precision"),
                            value(row, "macro_recall"), value(row, "micro_precision"), value(row, "micro_recall")])
    parts += ["**BẢNG IV. PRECISION/RECALL CÙNG NGƯỠNG 0,5.**",
              table(["Hệ", "Split", "P macro", "R macro", "P micro", "R micro"], pr_rows),
              "Split được ghi theo từng hàng. Không xếp hạng một hàng val với một hàng test như cùng phép đo."]

    lookup = {(row["system"], row.get("seed"), row["split"], row["threshold_mode"]): row
              for row in summary.get("records", [])}
    seed_rows = []
    for run in data["runs"]:
        system = "C_" + run["architecture"]
        fixed = lookup.get((system, run["seed"], "test", "fixed"), {})
        tuned = lookup.get((system, run["seed"], "test", "tuned"), {})
        seed_rows.append([SHORT[system], run["seed"], run["metadata"].get("selected_epoch", "—") if run["completed"] else "—",
                          fmt(run["metrics"].get("macro_f1")), fmt(fixed.get("macro_f1")), fmt(tuned.get("macro_f1"))])
    parts += ["**BẢNG V. TỪNG SEED C; MACRO-F1.**",
              table(["C", "Seed", "Epoch", "Val fixed", "Test fixed", "Test tuned"], seed_rows),
              "Epoch chọn bằng validation @0,5 của chính seed. Bảy metrics từng seed/split/ngưỡng nằm trong "
              "reports/project_results/all_runs.csv; mean_std.csv giữ toàn bộ nhóm, không chỉ bảng rút gọn."]

    b = data["zero_shot"]
    if b and b["mode"] == "smoke":
        bm = b["metrics"]
        parts += [f"**B pilot, ngoài bảng benchmark:** mới đo {bm.get('n_samples', '—')} câu, "
                  f"Macro-F1 {fmt(bm.get('macro_f1'))}, Micro-F1 {fmt(bm.get('micro_f1'))}, "
                  f"Hamming {fmt(bm.get('hamming_loss'))}. Số này chỉ kiểm pipeline; "
                  "không dùng để kết luận B tốt/kém hơn các hàng full."]
    if summary.get("missing"):
        parts += ["Các phần chưa đủ trong summary: " + "; ".join(summary["missing"]) + "."]
    parts += ["Hồ sơ đối chiếu gồm summary.json (revision, sources và hashes), all_runs.csv, mean_std.csv, "
              "per_label.csv, final_protocol.json của từng run và reports/reproducibility. "
              "SHA-256 summary của bản xuất này lưu trong BAI_BAO_GOEMOTIONS_IEEE_KIEM_CHUNG.json. "
              "Trọng số lớn/checkpoint lưu riêng trong data/processed, không coi việc vắng chúng trên Git là thiếu mô tả phương pháp."]
    return "\n\n".join(parts)


def rare(data):
    rows = read_csv(ROOT / "reports/project_results/rare_before_after.csv")
    manifest_path = ROOT / "reports/project_results/analysis_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    primary = "test" if manifest.get("rare_test_complete") else "validation"
    chosen = [row for row in rows if row.get("split") == primary]
    parts = [f"**BẢNG VI. NĂM NHÃN HIẾM TRƯỚC/SAU — {'TEST' if primary == 'test' else 'VALIDATION CÓ CALIBRATION'}.**"]
    if not chosen:
        return parts[0] + "\n\nChưa có hồ sơ cặp trước/sau hợp lệ. Không chọn nhãn hiếm bằng test hoặc tạo số thay thế."
    tab = []
    for row in sorted(chosen, key=lambda item: (int(float(item.get("train_support") or 0)), item.get("method", ""), item.get("architecture", ""))):
        good = row.get("status") == "ok"
        system = "A" if row.get("method") == "A" else SHORT.get("C_" + row.get("architecture", ""), "C")
        tab.append([row.get("label", "—"), system, row.get("train_support", "—"),
                    str(int(float(row["evaluation_support"]))) if good and number(row.get("evaluation_support")) is not None else "—",
                    fmt(row.get("before_f1_mean"), row.get("before_f1_std")) if good else "—",
                    fmt(row.get("after_f1_mean"), row.get("after_f1_std")) if good else "—",
                    fmt(row.get("delta_f1_mean"), sign=True) if good else "—"])
    parts += [table(["Nhãn", "Hệ", "Train+", "Eval+", "F1 trước", "F1 sau", "Δ"], tab),
              "Train+/Eval+ là support dương trong train/split đánh giá. A: standard fixed→balanced tuned, một run; "
              "C: fixed→tuned cùng checkpoint/seed, đủ ba seed mới có mean±std. Δ là trung bình chênh lệch cặp seed; "
              "mọi tăng/giảm và std của Δ được giữ trong rare_before_after.csv. Ô — là thiếu/chưa đủ dữ liệu."]
    if primary == "validation":
        parts += ["Bảng hiện tại dùng validation đã hiệu chỉnh ngưỡng; test trước/sau chưa đủ. "
                  "Không gọi các mức tăng này là cải thiện test hay chứng cứ triển khai thực tế."]
    return "\n\n".join(parts)


def errors():
    parts = ["**Ví dụ A standard @0,5 trên validation đã kiểm.** ID eczwil0, “I am so proud of this community.”, "
             "nhãn thật pride nhưng dự đoán rỗng, score pride=0,3776; ID ed832y6, “Homeopaths love it!”, "
             "nhãn thật neutral nhưng dự đoán love, score love≈1; ID eczdvun, “Thank you. I really appreciate your response”, "
             "nhãn thật admiration/gratitude nhưng chỉ tìm gratitude, score admiration=0,4989. "
             "Đây là lỗi A; không gán các ví dụ này thành lỗi của C khi chưa đọc scores C."]
    folder = ROOT / "reports/errors_test_standard_fixed"
    counts, examples = read_csv(folder / "counts.csv"), read_csv(folder / "examples.csv")
    if not counts or not examples:
        return (parts[0] + "\n\n**C1/C2/C3 chưa có hồ sơ so sánh lỗi test đầy đủ.** Cần ba checkpoint đại diện "
                "chọn bằng validation, cùng ID và ít nhất ba nhóm. Có code phân tích chưa chứng minh đã hoàn tất số đo.")
    tab = [[SHORT.get("C_" + row["architecture"], row["architecture"]), row["seed"],
            row["category"].replace("partial_multi_label", "Thiếu nhãn").replace("rare_false_negative", "Bỏ nhãn hiếm").replace("missed_extra_pair", "FN+FP"),
            row["error_samples"], row["eligible_samples"], f"{float(row['fraction_eligible']):.2%}"] for row in counts]
    parts += ["**BẢNG VII. BA NHÓM LỖI C TRÊN CÙNG TEST, NGƯỠNG 0,5.**",
              table(["C", "Seed", "Nhóm", "Lỗi", "Đủ điều kiện", "Tỷ lệ"], tab),
              "Các seed đại diện chọn theo validation của từng C. Đếm lỗi này không phải mean±std qua ba seed; "
              "nhóm có thể chồng lấp. Manifest lưu mapping, nguồn và hash."]
    for category in ("partial_multi_label", "rare_false_negative", "missed_extra_pair"):
        ids = sorted({row["id"] for row in examples if row["category"] == category})
        if not ids:
            parts += [f"Nhóm {category}: không có ví dụ đáp ứng trong hồ sơ; giữ trạng thái này."]
            continue
        chosen = [row for row in examples if row["category"] == category and row["id"] == ids[0]]
        parts += [f"**{category}, ID {ids[0]}.** Văn bản: “{chosen[0]['text']}”. Nhãn thật: "
                  + ", ".join(json.loads(chosen[0]["true_labels"])) + "."]
        for row in chosen:
            predicted = ", ".join(json.loads(row["predicted_labels"])) or "không nhãn"
            missed = ", ".join(json.loads(row["missed_labels"])) or "không"
            extra = ", ".join(json.loads(row["extra_labels"])) or "không"
            parts += [f"{SHORT.get('C_' + row['architecture'], row['architecture'])} seed {row['seed']}: "
                      f"dự đoán {predicted}; bỏ sót {missed}; nhãn thừa {extra}; "
                      f"gặp nhóm lỗi đang xét: {row['error_present']}. "
                      + (f"Nhận xét đã điền: {row['manual_linguistic_notes']}" if row.get("manual_linguistic_notes") else "Nhận xét ngôn ngữ cần nhóm đọc thủ công.")]
    parts += ["Ví dụ chọn theo ID có thứ tự từ union các model, cùng ID cho cả ba C; không chọn riêng những câu thuận lợi cho một mô hình. "
              "Đầy đủ điểm 28 nhãn, các ví dụ còn lại và cặp nhầm nằm trong examples.csv/pairs.csv."]
    return "\n\n".join(parts)


def discussion(data):
    averages = aggregate_lookup(data)
    candidates = [averages.get((name, "validation", "fixed")) for name in ("C_bert", "C_roberta", "C_distilbert")]
    candidates = [row for row in candidates if row and row.get("n_runs") == 3]
    parts = []
    if len(candidates) == 3:
        ordered = sorted(candidates, key=lambda row: -row["macro_f1_mean"])
        stable = min(candidates, key=lambda row: row["macro_f1_std"])
        parts += [f"Trong ba cấu hình đã thử, {SHORT[ordered[0]['system']]} đạt mean Macro-F1 validation @0,5 cao nhất "
                  f"({value(ordered[0], 'macro_f1')}); {SHORT[ordered[-1]['system']]} thấp nhất "
                  f"({value(ordered[-1], 'macro_f1')}). {SHORT[stable['system']]} có sample std Macro-F1 nhỏ nhất "
                  f"({fmt(stable['macro_f1_std'])}). Đây là thứ hạng các hệ được triển khai với cấu hình đã ghi, "
                  "không khẳng định một kiến trúc luôn tốt nhất."]
    else:
        parts += ["Chưa đủ ba seed của cả ba C nên chưa xếp hạng hoặc khẳng định C1/C2/C3 tốt nhất/ổn định nhất. "
                  "Một run full giúp kiểm triển khai, nhưng không thay điều kiện so sánh mean±std. "
                  "A-W tuned-validation không được so với một C fixed-validation rồi quy ưu thế hoàn toàn cho mô hình."]
    for system in ("C_bert", "C_roberta", "C_distilbert"):
        before, after = averages.get((system, "test", "fixed")), averages.get((system, "test", "tuned"))
        if before and after:
            parts += [f"{SHORT[system]} test Macro-F1 fixed→tuned: {value(before, 'macro_f1')}→{value(after, 'macro_f1')}; "
                      f"Micro-F1 {value(before, 'micro_f1')}→{value(after, 'micro_f1')}. "
                      "Đối chiếu support/P/R/Hamming và mức giảm ở từng nhãn trước khi gọi cải thiện tổng thể."]
    selection = data["selection"]
    if selection:
        parts += [f"Demo C được chọn theo hồ sơ: {selection.get('architecture', '—')}, seed {selection.get('seed', '—')}; "
                  f"rule kiến trúc: {selection.get('selection_rule', '—')}; rule checkpoint: {selection.get('checkpoint_rule', '—')}; "
                  f"ngưỡng mặc định: {selection.get('default_threshold', '—')}. Nguồn selected_model.json; không chọn lại bằng test."]
    else:
        parts += ["Chưa có selected_model.json hợp lệ cho best C ở thời điểm xuất; chưa xác nhận demo đạt yêu cầu D."]
    demo = data["demo"]
    ui_path = ROOT / "reports/demo_ui/evidence.json"
    ui = json.loads(ui_path.read_text(encoding="utf-8")) if ui_path.exists() else {}
    if demo:
        parts += ["Đã có hồ sơ kiểm suy luận demo (reports/demo_verification.json). Trạng thái ghi nhận: "
                  + str(demo.get("status", demo.get("passed", demo.get("checks", "xem hồ sơ"))))
                  + "; đây là kiểm tại thời điểm hồ sơ, không tự khẳng định server hiện đang mở."]
    else:
        parts += ["Chưa có hồ sơ verify_demo hoàn tất tại thời điểm xuất. Giao diện cần kiểm với checkpoint thật."]
    if ui:
        parts += [f"Kiểm giao diện lúc {ui.get('checked_at_utc', '—')}: {ui.get('interface_status', '—')}, "
                  f"{ui.get('rendered_row_count', '—')}/28 hàng score; tương đương điểm model ở kiểm UI: "
                  f"{ui.get('model_score_equivalence', '—')}. Kiểm giao diện và kiểm suy luận là hai bằng chứng riêng."]
    return "\n\n".join(parts)


def conclusion(data):
    completed = sum(row["completed"] for row in data["runs"])
    complete = data["summary"].get("complete", False)
    text = ("Đồ án giữ bài toán 28 nhãn và official split của GoEmotions, xây A dễ giải thích, "
            "B không fine-tune và ba C fine-tune, cùng module đánh giá/threshold. "
            "Class weighting và ngưỡng riêng được đối chiếu bằng ablation, support nhãn hiếm và các ví dụ lỗi. ")
    if complete:
        text += ("Đã có A/B/C full cùng bảng validation/test; C gồm đủ ba seed mỗi kiến trúc. "
                 "Test sử dụng mô hình/ngưỡng đã khóa trên validation. Demo và hồ sơ đối chiếu vẫn cần nhóm kiểm khi chuyển máy. ")
    else:
        text += (f"Hiện C hoàn tất {completed}/9 run full; phần chưa có B/C/test/lỗi/demo được đánh dấu trong bảng. "
                 "Không dùng pilot hoặc điểm tuned-validation để thay kết luận test. "
                 "Cần hoàn tất các artifact thiếu, khóa protocol rồi đánh giá test và kiểm demo trước khi dùng bản nộp cuối. ")
    return text + ("Giá trị đầu ra là cung cấp điểm và nhãn hỗ trợ đọc/phân tích phản hồi. "
                   "Hiệu quả tại domain doanh nghiệp, tiếng Việt và ROI cần dữ liệu cùng quy trình đo riêng. "
                   "Liên hệ Case Study 4 giúp giải thích mối quan hệ dữ liệu→quyết định→giá trị, "
                   "không biến ví dụ tiết kiệm năng lượng thành bằng chứng tài chính của NLP.")


def refresh(source, data):
    stamp = datetime.now(timezone(timedelta(hours=7))).strftime("%d/%m/%Y %H:%M (UTC+7)")
    content = source.read_text(encoding="utf-8")
    refs = {row["id"]: row for row in json.loads((ROOT / "reports/references_ieee.json").read_text(encoding="utf-8"))}
    blocks = {"ABSTRACT": abstract(data), "RESULTS": results(data, stamp), "RARE": rare(data), "ERRORS": errors(),
              "DISCUSSION": discussion(data), "CONCLUSION": conclusion(data),
              "REFERENCES": "\n\n".join(f"[{index}] {refs[original]['ieee']}" for index, original in enumerate(REF_IDS, 1))}
    for name, replacement in blocks.items():
        pattern = rf"<!-- AUTO_{name} -->.*?<!-- END_AUTO_{name} -->"
        content, count = re.subn(pattern, lambda _: f"<!-- AUTO_{name} -->\n{replacement}\n<!-- END_AUTO_{name} -->", content, flags=re.DOTALL)
        if count != 1:
            raise ValueError(f"Cần đúng một khối AUTO_{name}, thấy {count}")
    # Kiểm thứ tự trích dẫn lần đầu, tách bibliography khỏi nội dung.
    body = content.split("## TÀI LIỆU THAM KHẢO")[0]
    first_order = list(dict.fromkeys(int(x) for x in re.findall(r"\[(\d+)\]", body)))
    if first_order != list(range(1, len(REF_IDS) + 1)):
        raise ValueError(f"Thứ tự trích dẫn IEEE không liên tiếp: {first_order}")
    source.write_text(content, encoding="utf-8")
    return content, stamp


def page_layout(section, *, columns):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(1.0)
    section.left_margin = section.right_margin = Inches(0.625)
    section.header_distance = section.footer_distance = Inches(0.3)
    col = section._sectPr.find(qn("w:cols"))
    if col is None:
        col = OxmlElement("w:cols")
        section._sectPr.append(col)
    col.set(qn("w:num"), str(columns))
    col.set(qn("w:space"), "360")  # 0.25 inch = 360 twips.
    col.set(qn("w:equalWidth"), "1")
    # Theo mẫu conference: không tạo header/footer/PAGE hoặc mã copyright giả.


def styles(document):
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.35)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Normal", 10), ("Heading 1", 10), ("Heading 2", 10), ("Heading 3", 10), ("Caption", 8)):
        style = document.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "Times New Roman")
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        heading = document.styles[name]
        heading.font.bold = False
        heading.font.italic = name != "Heading 1"
        heading.paragraph_format.page_break_before = False
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.first_line_indent = Cm(0)
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(3)
        heading.paragraph_format.line_spacing = 1.0
        heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
    document.core_properties.title = "Phân loại cảm xúc đa nhãn với GoEmotions"
    document.core_properties.subject = "Bài đồ án theo định dạng IEEE conference hai cột"
    document.core_properties.author = "Bảo Duy Nguyễn; Quốc Khánh; Đức Trí; Nhật Huy"
    document.core_properties.keywords = "GoEmotions, NLP, multi-label, IEEE conference"


def table_word(document, rows, *, border=True, widths=None, font=8):
    result = document.add_table(rows=1, cols=len(rows[0]))
    result.alignment, result.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    result.style = "Table Grid" if border else "Normal Table"
    widths = widths or [COLUMN_CM / len(rows[0])] * len(rows[0])
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    for index, row in enumerate(rows):
        cells = result.rows[0].cells if index == 0 else result.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width, cell.vertical_alignment = Cm(width), WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for direction in ("top", "left", "bottom", "right"):
                element = OxmlElement("w:" + direction)
                element.set(qn("w:w"), "35")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            cell._tc.get_or_add_tcPr().append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = index == 0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(str(text)) > 16 else WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, str(text), size=font)
            for run in paragraph.runs:
                run.bold = border and index == 0
        no_split = OxmlElement("w:cantSplit")
        result.rows[index]._tr.get_or_add_trPr().append(no_split)
        if index == 0 and border:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            result.rows[0]._tr.get_or_add_trPr().append(repeat)
    return result


def reference_word(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(-0.40)
    paragraph.paragraph_format.left_indent = Cm(0.40)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for part in re.split(r"(https?://[^\s]+)", text):
        if re.match(r"https?://", part):
            hyperlink(paragraph, part, part)
        else:
            set_font(paragraph.add_run(part), size=8)
    # Hyperlink helper dùng font paragraph, ép size8 cho cả hyperlink runs.
    for run in paragraph._p.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(0, properties)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "16")
        properties.append(size)


def build(content, output):
    document = Document()
    styles(document)
    page_layout(document.sections[0], columns=1)
    title = content.splitlines()[0].removeprefix("# ")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(12)
    set_font(paragraph.add_run(title), size=24)
    paragraph = document.add_paragraph("Bảo Duy Nguyễn    Quốc Khánh    Đức Trí    Nhật Huy")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        set_font(run, size=11)
    paragraph = document.add_paragraph("Đơn vị: ____________________    Email: ____________________")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        set_font(run, size=10)
    # Title/authors toàn chiều rộng; abstract và phần còn lại theo hai cột.
    page_layout(document.add_section(WD_SECTION_START.CONTINUOUS), columns=2)
    lines, index, references = content.splitlines()[1:], 0, False
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("<!--"):
            index += 1
            continue
        if line.startswith("## "):
            text = line[3:]
            references = text == "TÀI LIỆU THAM KHẢO"
            document.add_paragraph(text, style="Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif references and re.match(r"^\[\d+\] ", line):
            reference_word(document, line)
        elif line.startswith("| "):
            group = []
            while index < len(lines) and lines[index].strip().startswith("| "):
                group.append(lines[index].strip())
                index += 1
            table_word(document, table_rows(group))
            continue
        elif line.startswith("$$ ") and line.endswith(" $$"):
            text = line[3:-3]
            match = re.match(r"(.+?)\s+(\(\d+\))$", text)
            formula, label = (match[1], match[2]) if match else (text, "")
            eq = table_word(document, [[formula, label]], border=False, widths=[COLUMN_CM - 0.5, 0.5], font=9)
            eq.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            eq.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif re.match(r"^\*\*BẢNG [IVX]+\.", line):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            inline(paragraph, line, size=8)
        else:
            paragraph = document.add_paragraph()
            abstract_or_keywords = line.startswith(("**Tóm tắt", "**Từ khóa"))
            inline(paragraph, line, size=9 if abstract_or_keywords else 10)
            if abstract_or_keywords:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(6)
                for run in paragraph.runs:
                    run.bold = True
                    run.italic = line.startswith("**Từ khóa")
        index += 1
    # Continuous final section balances the last pair of columns in Word.
    page_layout(document.add_section(WD_SECTION_START.CONTINUOUS), columns=2)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return document


def available_memory_mb():
    if sys.platform != "win32":
        return None
    class MemoryStatus(ctypes.Structure):
        _fields_ = [("length", ctypes.c_ulong), ("load", ctypes.c_ulong),
                    *( (name, ctypes.c_ulonglong) for name in ("total", "available", "page_total", "page_available", "virtual_total", "virtual_available", "extended") )]
    status = MemoryStatus()
    status.length = ctypes.sizeof(status)
    if not ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(status)):
        return None
    return status.available / (1024 * 1024)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--pdf", action="store_true", help="Xuất PDF bằng Word sẵn có khi RAM trống >=700MB")
    args = parser.parse_args()
    data = snapshot()
    content, stamp = refresh(args.source, data)
    document = build(content, args.output)
    evidence = {"updated": stamp, "format": "IEEE conference two columns; Vietnamese course manuscript, not an IEEE publication",
                "format_sources": FORMAT_SOURCES,
                "paper_inches": [8.5, 11], "margins_inches": {"top": 0.75, "bottom": 1, "left": 0.625, "right": 0.625},
                "columns": 2, "column_width_inches": 3.5, "gap_inches": 0.25,
                "fonts_points": {"title": 24, "authors": 11, "body": 10, "abstract": 9, "references": 8, "tables": 8},
                "known_authors": document.core_properties.author,
                "unknown_admin_fields": "blank", "reference_count": len(REF_IDS),
                "first_citation_order": list(range(1, len(REF_IDS) + 1)),
                "c_completed": sum(row["completed"] for row in data["runs"]), "c_expected": 9,
                "records": len(data["summary"].get("records", [])), "average_groups": len(data["summary"].get("averages", [])),
                "summary_complete": data["summary"].get("complete", False), "summary_sha256": data["summary_hash"],
                "source_sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
                "docx": str(args.output), "pdf": None, "pdf_status": "not_requested", "no_model_import_or_training": True}
    if args.pdf:
        memory = available_memory_mb()
        evidence["free_memory_mb_before_word"] = memory
        if memory is not None and memory < 700:
            evidence["pdf_status"] = "deferred_low_memory"
            print(f"Đã xuất DOCX. RAM trống {memory:.0f}MB: giữ PDF chờ; chạy lại --pdf sau khi huấn luyện xong.")
        else:
            pdf = export_pdf_with_word(args.output)
            evidence.update(pdf=str(pdf), pdf_status="exported")
            print("PDF:", pdf)
    evidence_path = args.output.with_name(args.output.stem + "_KIEM_CHUNG.json")
    evidence_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print("DOCX:", args.output)
    print("Nguồn:", args.source)
    print("Đối chiếu:", evidence_path)
    print(f"{len(content.split()):,} từ tách khoảng trắng; {len(document.tables)} bảng gồm công thức; {len(REF_IDS)} nguồn.")


if __name__ == "__main__":
    main()
