"""Xuất báo cáo Word theo mẫu học phần từ bản Markdown đã kiểm nguồn.

Chạy ở thư mục repository:
    python tools/build_report_docx.py

Chỉ cần python-docx. Nội dung và số thực nghiệm lấy từ Markdown; script
không tạo số liệu hay thay đổi báo cáo nguồn. Có thể chạy lại sau khi bổ
sung kết quả. Mục lục/danh mục dùng Word fields và cập nhật khi mở Word.
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "reports" / "BAO_CAO_DO_AN_NOI_DUNG.md"
DEFAULT_OUTPUT = ROOT / "reports" / "BAO_CAO_DO_AN_GOEMOTIONS_IEEE.docx"
BODY_FONT = "Times New Roman"
CONTENT_WIDTH_CM = 21.59 - 2.5 - 2.0


def set_font(run, name=BODY_FONT, size=12):
    """Gán cả font Latin/Unicode để tiếng Việt hiển thị nhất quán."""
    run.font.name = name
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def field(paragraph, instruction, placeholder="", hidden=False):
    """Tạo field Word: PAGE, TOC hoặc TC mà không giả lập số trang."""
    run = paragraph.add_run()
    set_font(run)
    run.font.hidden = hidden
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    run._r.append(code)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    if placeholder:
        text = OxmlElement("w:t")
        text.text = placeholder
        run._r.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    return run


def hyperlink(paragraph, label, url):
    """Chèn liên kết thật; để URL nguyên văn trong danh mục IEEE."""
    from docx.opc.constants import RELATIONSHIP_TYPE

    relation = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "154B73")
    props.append(color)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    props.append(fonts)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))")


def inline(paragraph, text, size=12):
    """Đọc một phần nhỏ Markdown cần dùng: đậm, inline code, link."""
    for part in INLINE.split(text):
        if not part:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", part)
        if match:
            hyperlink(paragraph, match[1], match[2])
        else:
            run = paragraph.add_run(part[2:-2] if part.startswith("**") else part[1:-1] if part.startswith("`") else part)
            set_font(run, "Consolas" if part.startswith("`") else BODY_FONT, size)
            if part.startswith("**"):
                run.bold = True


def plain(text):
    return text.replace("**", "").replace("`", "")


def set_page(section, numbering=None, start=None):
    """Mẫu cô dùng Letter, lề trái 2,5 cm; ba lề còn lại 2 cm."""
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    if numbering:
        numbers = OxmlElement("w:pgNumType")
        numbers.set(qn("w:fmt"), numbering)
        if start is not None:
            numbers.set(qn("w:start"), str(start))
        section._sectPr.append(numbers)


def page_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    field(p, " PAGE ")
    for run in p.runs:
        set_font(run, size=10)


def configure(document):
    def style_font(style, size):
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        # Xóa theme font của mẫu python-docx (Calibri/Cambria) để Word
        # không thay Times New Roman của các tiêu đề khi cập nhật fields.
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
            if qn(f"w:{attr}") in fonts.attrib:
                del fonts.attrib[qn(f"w:{attr}")]
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), BODY_FONT)

    normal = document.styles["Normal"]
    style_font(normal, 12)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13), ("Heading 4", 12)):
        style = document.styles[name]
        style_font(style, size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    style_font(document.styles["Caption"], 11)
    document.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Caption"].paragraph_format.space_before = Pt(4)
    document.styles["Caption"].paragraph_format.space_after = Pt(6)
    document.styles["Caption"].paragraph_format.line_spacing = 1.1
    document.styles["Caption"].paragraph_format.keep_with_next = True
    for name in ("FigureCaption", "TableCaption"):
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Caption"]
        style_font(style, 11)
    for name in ("List Bullet", "List Number"):
        style_font(document.styles[name], 12)
    for level in (1, 2, 3):
        name = f"TOC {level}"
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style_font(style, 11)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.space_after = Pt(4)
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.core_properties.title = "Phân loại cảm xúc đa nhãn với GoEmotions"
    document.core_properties.subject = "Báo cáo đồ án môn Xử lý ngôn ngữ tự nhiên"
    document.core_properties.author = "Bảo Duy Nguyễn; Quốc Khánh; Đức Trí; Nhật Huy"
    document.core_properties.keywords = "GoEmotions, NLP, multi-label, TF-IDF, BERT, IEEE"


def cover(document):
    def centered(text, size=14, bold=False, before=0, after=8):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        set_font(r, size=size)
        r.bold = bold
        return p

    centered("TRƯỜNG ĐẠI HỌC NGOẠI NGỮ – TIN HỌC\nTHÀNH PHỐ HỒ CHÍ MINH", 14, True)
    centered("KHOA CÔNG NGHỆ THÔNG TIN", 14, True)
    centered("BÁO CÁO ĐỒ ÁN MÔN HỌC", 16, True, before=36)
    centered("XỬ LÝ NGÔN NGỮ TỰ NHIÊN", 16, True)
    centered("PHÂN LOẠI CẢM XÚC ĐA NHÃN\nTRÊN VĂN BẢN MẠNG XÃ HỘI\nVỚI GOEMOTIONS", 18, True, before=28, after=24)
    for label in ("Giảng viên hướng dẫn", "Mã lớp học phần", "Năm học / học kỳ"):
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.1)
        inline(p, f"**{label}:** ____________________")
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.1)
    inline(p, "**Sinh viên thực hiện:**")
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(8)
    table.columns[1].width = Cm(5)
    for c, text in zip(table.rows[0].cells, ("Họ và tên", "Mã số sinh viên")):
        inline(c.paragraphs[0], f"**{text}**")
    for name in ("Bảo Duy Nguyễn", "Quốc Khánh", "Đức Trí", "Nhật Huy"):
        cells = table.add_row().cells
        inline(cells[0].paragraphs[0], name)
        inline(cells[1].paragraphs[0], "____________________")
    centered("Thành phố Hồ Chí Minh", 12, before=24)


def contents(document):
    document.add_heading("MỤC LỤC", 1)
    p = document.add_paragraph()
    field(p, ' TOC \\o "1-3" \\h \\z \\u ', "Cập nhật mục lục trong Word: Ctrl+A, F9.")
    document.add_heading("DANH MỤC HÌNH", 1)
    field(document.add_paragraph(), ' TOC \\t "FigureCaption,1" \\h \\z ', "Cập nhật danh mục hình trong Word.")
    document.add_heading("DANH MỤC BẢNG", 1)
    field(document.add_paragraph(), ' TOC \\t "TableCaption,1" \\h \\z ', "Cập nhật danh mục bảng trong Word.")


def caption(document, text):
    p = document.add_paragraph(style="FigureCaption" if text.startswith("Hình") else "TableCaption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if text.startswith("Hình") else WD_ALIGN_PARAGRAPH.LEFT
    inline(p, text, 11)
    return p


def table_rows(lines):
    result = []
    for line in lines:
        cells = [x.strip() for x in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
        if all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in cells):
            continue
        result.append(cells)
    return result


def set_cell_margins(cell, amount=65):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def add_table(document, rows):
    if not rows:
        return
    count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    weights = []
    for col in range(count):
        lengths = [len(plain(row[col])) for row in rows if col < len(row)]
        typical = max(10, min(44, sum(lengths) / len(lengths)))
        weights.append(typical ** 0.65)
    for col, weight in zip(table.columns, weights):
        col.width = Cm(CONTENT_WIDTH_CM * weight / sum(weights))
    font_size = 9 if count >= 7 else 10.5
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for col, cell in enumerate(cells):
            cell.width = table.columns[col].width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            inline(p, row[col] if col < len(row) else "", font_size)
            if index == 0:
                for run in p.runs:
                    run.bold = True
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "E6EDF3")
                cell._tc.get_or_add_tcPr().append(fill)
        props = table.rows[-1]._tr.get_or_add_trPr()
        avoid_split = OxmlElement("w:cantSplit")
        props.append(avoid_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            props.append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_reference_table(document, entries):
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(1)
    table.columns[1].width = Cm(CONTENT_WIDTH_CM - 1)
    for number, text in entries:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, 0)
        cells[0].width = Cm(1)
        cells[1].width = Cm(CONTENT_WIDTH_CM - 1)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(8)
        inline(p, f"[{number}]", 11)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        parts = re.split(r"(https?://\S+)", text)
        for part in parts:
            if part.startswith(("https://", "http://")):
                hyperlink(p, part, part)
            else:
                inline(p, part, 11)
        props = table.rows[-1]._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))


def build(source: Path, output: Path):
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()
    configure(document)
    set_page(document.sections[0])
    cover(document)
    front = document.add_section(WD_SECTION_START.NEW_PAGE)
    set_page(front, "lowerRoman", 1)
    page_footer(front)
    try:
        start = next(i for i, line in enumerate(lines) if line.strip() == "## Lời cảm ơn")
    except StopIteration:
        start = 0
    body_started = False
    first_front_heading = True
    in_references = False
    i = start
    image_count = table_count = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("<!--"):
            i += 1
            continue
        if line == "## Mục lục và danh mục":
            contents(document)
            i += 1
            while i < len(lines) and not lines[i].startswith("# "):
                i += 1
            continue
        if line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            p = document.add_paragraph()
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "F2F3F4")
            p._p.get_or_add_pPr().append(fill)
            run = p.add_run("\n".join(code_lines))
            set_font(run, "Consolas", 9)
            i += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading[1])
            title = plain(heading[2])
            if title.startswith("CHƯƠNG 1.") and not body_started:
                section = document.add_section(WD_SECTION_START.NEW_PAGE)
                set_page(section, "decimal", 1)
                page_footer(section)
                body_started = True
            in_references = title == "TÀI LIỆU THAM KHẢO"
            if not body_started:
                p = document.add_heading(title.upper(), 1)
                if first_front_heading:
                    p.paragraph_format.page_break_before = False
                    first_front_heading = False
            else:
                p = document.add_heading(title, level)
                if title.startswith("CHƯƠNG 1."):
                    p.paragraph_format.page_break_before = False
            i += 1
            continue
        ref = re.match(r"^\[(\d+)\]\s+(.+)$", line)
        if in_references and ref:
            entries = []
            while i < len(lines):
                value = re.match(r"^\[(\d+)\]\s+(.+)$", lines[i].strip())
                if value:
                    entries.append((value[1], value[2]))
                    i += 1
                elif not lines[i].strip():
                    i += 1
                else:
                    break
            add_reference_table(document, entries)
            table_count += 1
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^\)]+)\)", line)
        if image:
            path = Path(image[2].strip("<>"))
            if not path.is_absolute():
                path = source.parent / path
            if not path.exists():
                raise FileNotFoundError(f"Không tìm thấy ảnh báo cáo: {path}")
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm(16))
            image_count += 1
            i += 1
            continue
        title = plain(line)
        if re.match(r"^(Hình\s+\d+[.\-]\d+|Bảng\s+(?:\d+|[A-Z])[.\-]\d+[a-z]?)\.", title):
            caption(document, title)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(document, table_rows(rows))
            table_count += 1
            continue
        if re.match(r"^-\s+", line):
            p = document.add_paragraph(style="List Bullet")
            inline(p, line[2:])
        elif re.match(r"^\d+\.\s+", line):
            p = document.add_paragraph(style="List Number")
            inline(p, re.sub(r"^\d+\.\s+", "", line))
        else:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(p, line)
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {"output": str(output), "paragraphs": len(document.paragraphs), "tables": table_count, "images": image_count}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--markdown", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--pdf", action="store_true", help="Cập nhật fields và xuất PDF bằng Microsoft Word đã cài trên Windows")
    args = parser.parse_args()
    output = args.output.resolve()
    result = build(args.markdown.resolve(), output)
    if args.pdf:
        result["pdf"] = export_pdf_with_word(output)
    print(result)


def export_pdf_with_word(path):
    """Dùng Word có sẵn, chạy ẩn; chỉ đóng tài liệu script vừa mở.

    Không cài Word/COM packages. Nếu có tài liệu khác trong instance,
    script giữ application hoạt động để không đóng tài liệu của người dùng.
    """
    if os.name != "nt":
        raise RuntimeError("Xuất PDF bằng Word chỉ hỗ trợ Windows có Microsoft Word.")
    pdf = path.with_suffix(".pdf")
    source_literal = str(path).replace("'", "''")
    pdf_literal = str(pdf).replace("'", "''")
    script = f"""
$ErrorActionPreference = 'Stop'
$reportWordApp = $null
$reportWordDoc = $null
try {{
  $reportWordApp = New-Object -ComObject Word.Application
  $reportWordApp.Visible = $false
  $reportWordApp.DisplayAlerts = 0
  $reportWordDoc = $reportWordApp.Documents.Open('{source_literal}', $false, $false, $false)
  $reportWordDoc.Fields.Update() | Out-Null
  foreach ($reportToc in $reportWordDoc.TablesOfContents) {{ $reportToc.Update() }}
  $reportWordDoc.Repaginate()
  $reportWordDoc.Save()
  $reportWordDoc.ExportAsFixedFormat('{pdf_literal}', 17)
}} finally {{
  if ($null -ne $reportWordDoc) {{ $reportWordDoc.Close(0) }}
  if ($null -ne $reportWordApp -and $reportWordApp.Documents.Count -eq 0) {{ $reportWordApp.Quit() }}
}}
"""
    command = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True,
        timeout=120,
        creationflags=subprocess.CREATE_NO_WINDOW,
    )
    if command.returncode:
        diagnostic = command.stderr.decode("utf-8", errors="replace")
        raise RuntimeError(f"Word chưa xuất được PDF; DOCX được giữ nguyên. {diagnostic}")
    if not pdf.exists():
        raise RuntimeError("Word không trả lỗi nhưng chưa tạo PDF; kiểm lại ứng dụng Word.")
    return str(pdf)


if __name__ == "__main__":
    main()
